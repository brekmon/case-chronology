#!/usr/bin/env python3
"""
make_samples.py - generate a synthetic document set to test against.

    python samples/make_samples.py

WHY THE SAMPLES ARE GENERATED RATHER THAN COMMITTED
----------------------------------------------------
Every document this produces is invented. The parties, the dates, the figures
and the matter itself do not exist. Generating them from a script rather than
committing a folder of files means anyone reading this repository can verify
that for themselves in about thirty seconds, which is a stronger claim than a
sentence in a README promising it.

If you want to try the tools on real material, use published court opinions or
other genuine public record. Never use documents from a live matter, and never
use redacted ones: redaction failures are common and a repository is forever.

The sample matter is a commercial contract dispute. That is deliberate. The
tooling is domain-general, and a neutral matter type keeps the demonstration
about the software.

WHAT IT DELIBERATELY BREAKS
---------------------------
One exhibit is generated as an image with NO text layer, exactly like a page
that was scanned and never OCR'd. It looks like a normal page to a human and
like a blank page to a text extractor. coverage.py is supposed to catch it. If
you change these samples, keep something broken in them: a gate you have never
seen fail is a gate you have no reason to trust.
"""
import os

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = LETTER

CAPTION = [
    "IN THE DISTRICT COURT OF MERIDIAN COUNTY",
    "STATE OF FRANKLIN",
    "",
    "MERIDIAN SUPPLY CO.,",
    "                              Plaintiff,",
    "v.                                                    Case No. 2024-CV-00817",
    "NORTHGATE FABRICATION LLC,",
    "                              Defendant.",
    "",
    "All parties, dates and figures in these documents are fictional.",
]


def page_header(c, title):
    c.setFont("Helvetica-Bold", 11)
    y = H - 1.0 * inch
    for line in CAPTION:
        c.drawString(1.0 * inch, y, line)
        y -= 14
        c.setFont("Helvetica", 10)
    y -= 10
    c.setFont("Helvetica-Bold", 13)
    c.drawString(1.0 * inch, y, title)
    return y - 26


def body(c, y, lines):
    c.setFont("Helvetica", 10.5)
    for line in lines:
        if y < 1.1 * inch:
            c.showPage()
            y = H - 1.0 * inch
            c.setFont("Helvetica", 10.5)
        c.drawString(1.0 * inch, y, line[:100])
        y -= 15
    return y


def motion():
    p = os.path.join(HERE, "01_motion_to_compel.pdf")
    c = canvas.Canvas(p, pagesize=LETTER)
    y = page_header(c, "MOTION TO COMPEL PRODUCTION")
    y = body(c, y, [
        "1. On March 4, 2024, Plaintiff served its First Request for Production.",
        "2. Defendant's responses were due on April 3, 2024.",
        "3. No responses were received by April 3, 2024.",
        "4. Counsel conferred by telephone on April 11, 2024.",
        "5. Defendant produced 214 pages on April 19, 2024, omitting Exhibit C.",
        "6. Plaintiff requested the omitted exhibit on April 22, 2024.",
        "7. As of May 6, 2024, Exhibit C has not been produced.",
        "",
        "WHEREFORE Plaintiff respectfully requests an order compelling production",
        "of all documents responsive to Request No. 7 within fourteen days.",
        "",
        "Dated this 6th day of May, 2024.",
    ])
    c.showPage()
    y = page_header(c, "CERTIFICATE OF SERVICE")
    body(c, y, [
        "I certify that on May 6, 2024, a copy of the foregoing was served on",
        "counsel of record by electronic filing.",
    ])
    c.save()
    print("  wrote 01_motion_to_compel.pdf  (2 pages, text)")


def statement(version, total, rows):
    p = os.path.join(HERE, f"0{version + 1}_financial_statement_v{version}.pdf")
    c = canvas.Canvas(p, pagesize=LETTER)
    y = page_header(c, f"SWORN FINANCIAL STATEMENT (VERSION {version})")
    lines = ["ASSETS"] + [f"    {n:<44}{v:>14}" for n, v in rows] + [
        "", f"    {'TOTAL':<44}{total:>14}", "",
        "I declare under penalty of perjury that the foregoing is true and correct.",
    ]
    body(c, y, lines)
    c.save()
    print(f"  wrote 0{version + 1}_financial_statement_v{version}.pdf  (1 page, text)")


def scanned_exhibit():
    """A page that is an image and nothing else. This is the planted defect."""
    p = os.path.join(HERE, "04_exhibit_c_scanned.pdf")
    img_path = os.path.join(HERE, "_exhibit_c_page.png")

    im = Image.new("RGB", (1700, 2200), "white")
    d = ImageDraw.Draw(im)
    try:
        f_hd = ImageFont.truetype("times.ttf", 46)
        f_bd = ImageFont.truetype("times.ttf", 34)
    except Exception:
        f_hd = f_bd = ImageFont.load_default()
    d.text((150, 180), "EXHIBIT C", font=f_hd, fill="black")
    d.text((150, 300), "PURCHASE ORDER 4471", font=f_bd, fill="black")
    for i, line in enumerate([
        "Date: February 12, 2024",
        "Vendor: Northgate Fabrication LLC",
        "Quantity: 1,200 units",
        "Unit price: $18.40",
        "Total: $22,080.00",
        "",
        "This page exists only as an image.",
        "A text extractor sees nothing here.",
        "coverage.py is supposed to catch that.",
    ]):
        d.text((150, 420 + i * 60), line, font=f_bd, fill="black")
    d.rectangle([120, 140, 1580, 1000], outline="black", width=3)
    im.save(img_path)

    c = canvas.Canvas(p, pagesize=LETTER)
    c.drawImage(img_path, 0.6 * inch, 2.0 * inch,
                width=W - 1.2 * inch, height=(W - 1.2 * inch) * 2200 / 1700)
    c.save()
    os.remove(img_path)
    print("  wrote 04_exhibit_c_scanned.pdf  (1 page, IMAGE ONLY, no text layer)")


def correspondence():
    try:
        import docx
    except ImportError:
        print("  skipped correspondence.docx (python-docx not installed)")
        return
    p = os.path.join(HERE, "05_correspondence.docx")
    d = docx.Document()
    d.add_heading("Letter to Opposing Counsel", level=1)
    for para in [
        "April 22, 2024",
        "Re: Meridian Supply Co. v. Northgate Fabrication LLC, No. 2024-CV-00817",
        "Counsel,",
        "Your April 19 production omitted Exhibit C, referenced at page 6 of the "
        "purchase agreement. Please produce it by April 29, 2024.",
        "Regards,",
        "Counsel for Plaintiff",
    ]:
        d.add_paragraph(para)
    d.save(p)
    print("  wrote 05_correspondence.docx  (not paginated)")


def main():
    print(f"generating synthetic documents in {HERE}")
    motion()
    statement(1, "$412,500.00", [
        ("Operating account 8842", "$84,200.00"),
        ("Equipment, net", "$196,300.00"),
        ("Accounts receivable", "$132,000.00"),
    ])
    statement(2, "$389,100.00", [
        ("Operating account 8842", "$61,800.00"),
        ("Equipment, net", "$196,300.00"),
        ("Accounts receivable", "$131,000.00"),
    ])
    scanned_exhibit()
    correspondence()
    print("\n  Everything above is invented. No real matter, party or figure.")
    print("  Next: python ingest.py samples -o ingest.json")


if __name__ == "__main__":
    main()
